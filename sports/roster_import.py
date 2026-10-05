"""
Carga masiva de la plantilla (roster) de un equipo desde CSV / Excel.

Columnas obligatorias: Nombres, Apellidos, Documento de Identidad,
Número de Camiseta, Posición. Opcionales: Fecha de Nacimiento, Teléfono.
"""

from __future__ import annotations

import io
import unicodedata
from typing import Any

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from authentication.admin_import.parsing import (
    RowImportError,
    display_cell,
    parse_optional_datetime,
    parse_optional_int,
)
from ecommerce.batch_import import read_upload_rows

from .models import Player

ROSTER_HEADERS = [
    "Nombres",
    "Apellidos",
    "Documento de Identidad",
    "Número de Camiseta",
    "Posición",
    "Fecha de Nacimiento",
    "Teléfono",
]

REQUIRED_FIELDS = ("first_name", "last_name", "id_number", "jersey_number", "position")

FIELD_LABELS = {
    "first_name": "Nombres",
    "last_name": "Apellidos",
    "id_number": "Documento de Identidad",
    "jersey_number": "Número de Camiseta",
    "position": "Posición",
    "birth_date": "Fecha de Nacimiento",
    "phone": "Teléfono",
}

HEADER_ALIASES = {
    "first_name": {"nombres", "nombre", "first_name", "firstname"},
    "last_name": {"apellidos", "apellido", "last_name", "lastname"},
    "id_number": {
        "documento de identidad",
        "documento",
        "documento identidad",
        "cedula",
        "identificacion",
        "id_number",
    },
    "jersey_number": {
        "numero de camiseta",
        "numero camiseta",
        "camiseta",
        "dorsal",
        "numero",
        "jersey_number",
    },
    "position": {"posicion", "position"},
    "birth_date": {"fecha de nacimiento", "fecha nacimiento", "nacimiento", "birth_date"},
    "phone": {"telefono", "celular", "phone"},
}

POSITION_ALIASES = {
    "portero": "goalkeeper",
    "arquero": "goalkeeper",
    "defensa": "defender",
    "mediocampista": "midfielder",
    "volante": "midfielder",
    "delantero": "forward",
    "entrenador": "coach",
    "staff": "staff",
    "lanzador": "pitcher",
    "receptor": "catcher",
    "primera base": "first_base",
    "segunda base": "second_base",
    "tercera base": "third_base",
    "shortstop": "shortstop",
    "campocorto": "shortstop",
    "jardinero izquierdo": "left_field",
    "jardinero central": "center_field",
    "jardinero derecho": "right_field",
    "bateador designado": "designated_hitter",
    "utility": "utility",
    "p": "pitcher",
    "c": "catcher",
    "1b": "first_base",
    "2b": "second_base",
    "3b": "third_base",
    "ss": "shortstop",
    "lf": "left_field",
    "cf": "center_field",
    "rf": "right_field",
    "dh": "designated_hitter",
    "ut": "utility",
}

VALID_POSITIONS = {code for code, _label in Player.POSITION_CHOICES}
DB_BATCH_SIZE = 50


class RosterHeaderError(ValueError):
    def __init__(self, message: str, missing: list[str]):
        super().__init__(message)
        self.missing = missing


def _fold(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.lower().replace("_", " ").split())


_ALIAS_LOOKUP = {
    _fold(alias): field for field, aliases in HEADER_ALIASES.items() for alias in aliases
}


def map_headers(headers: list[str]) -> dict[str, str]:
    """Devuelve {campo_modelo: encabezado_original}. Lanza RosterHeaderError si faltan columnas."""
    mapping: dict[str, str] = {}
    for header in headers:
        field = _ALIAS_LOOKUP.get(_fold(header))
        if field and field not in mapping:
            mapping[field] = header
    missing = [FIELD_LABELS[f] for f in REQUIRED_FIELDS if f not in mapping]
    if missing:
        raise RosterHeaderError(
            "El formato de columnas no es el esperado. Faltan: "
            + ", ".join(missing)
            + ". Descarga la plantilla de ejemplo e inténtalo de nuevo.",
            missing,
        )
    return mapping


def normalize_position(value: Any) -> str | None:
    raw = _fold(value)
    if not raw:
        return None
    if raw.replace(" ", "_") in VALID_POSITIONS:
        return raw.replace(" ", "_")
    return POSITION_ALIASES.get(raw)


def _normalize_document(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return display_cell(value).replace(" ", "")


def build_roster_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "plantilla"
    ws.append(ROSTER_HEADERS)
    header_fill = PatternFill("solid", fgColor="1E293B")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    ws.append(["Juan", "Pérez", "1012345678", 10, "Delantero", "2001-05-14", "3001234567"])
    ws.append(["Carlos", "Gómez", "1098765432", 1, "Portero", "", ""])
    for col, width in zip("ABCDEFG", (18, 18, 24, 20, 18, 20, 16)):
        ws.column_dimensions[col].width = width

    help_ws = wb.create_sheet("posiciones")
    help_ws.append(["Posición (valor aceptado)", "Código"])
    for code, label in Player.POSITION_CHOICES:
        help_ws.append([label, code])
    help_ws.column_dimensions["A"].width = 28
    help_ws.column_dimensions["B"].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def roster_template_response() -> HttpResponse:
    response = HttpResponse(
        build_roster_template_xlsx(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_jugadores.xlsx"'
    return response


def _summary_message(created: int, error_count: int) -> str:
    noun = "jugador importado" if created == 1 else "jugadores importados"
    msg = f"{created} {noun} con éxito"
    if error_count:
        msg += f", {error_count} {'registro' if error_count == 1 else 'registros'} con errores"
    return msg


def import_team_roster(*, upload, team, user) -> dict[str, Any]:
    """
    Valida y crea los jugadores del archivo en el equipo.
    Lanza ValueError (formato/tamaño) o RosterHeaderError (columnas).
    """
    headers, rows = read_upload_rows(upload)
    mapping = map_headers(headers)
    if not rows:
        raise ValueError("El archivo no contiene filas de jugadores.")

    existing_docs = {
        _normalize_document(doc)
        for doc in team.players.exclude(id_number="").values_list("id_number", flat=True)
    }
    existing_jerseys = set(
        team.players.filter(is_active=True, jersey_number__isnull=False).values_list(
            "jersey_number", flat=True
        )
    )
    seen_docs: dict[str, int] = {}
    seen_jerseys: dict[int, int] = {}

    errors: list[dict[str, Any]] = []
    to_create: list[Player] = []

    def cell(row: dict, field: str) -> Any:
        header = mapping.get(field)
        return row.get(header, "") if header else ""

    for row in rows:
        r = int(row.get("_row") or 0)
        row_errors: list[tuple[str, str]] = []

        first_name = display_cell(cell(row, "first_name"))
        last_name = display_cell(cell(row, "last_name"))
        document = _normalize_document(cell(row, "id_number"))
        phone = display_cell(cell(row, "phone"))[:20]

        for field, value in (
            ("first_name", first_name),
            ("last_name", last_name),
            ("id_number", document),
        ):
            if not value:
                row_errors.append((field, f"'{FIELD_LABELS[field]}' es obligatorio."))

        jersey = None
        try:
            jersey = parse_optional_int(cell(row, "jersey_number"), FIELD_LABELS["jersey_number"])
            if jersey is None:
                row_errors.append(("jersey_number", "'Número de Camiseta' es obligatorio."))
            elif not 0 <= jersey <= 99:
                row_errors.append(("jersey_number", "El número de camiseta debe estar entre 0 y 99."))
                jersey = None
        except RowImportError as exc:
            row_errors.append(("jersey_number", str(exc)))

        raw_position = cell(row, "position")
        position = normalize_position(raw_position)
        if not display_cell(raw_position):
            row_errors.append(("position", "'Posición' es obligatoria."))
        elif position is None:
            row_errors.append(
                ("position", f"Posición '{display_cell(raw_position)}' no reconocida.")
            )

        birth_date = None
        try:
            parsed = parse_optional_datetime(
                cell(row, "birth_date"), FIELD_LABELS["birth_date"]
            )
            birth_date = parsed.date() if parsed else None
        except RowImportError as exc:
            row_errors.append(("birth_date", str(exc)))

        if document:
            if document in existing_docs:
                row_errors.append(
                    ("id_number", f"El documento {document} ya está registrado en este equipo.")
                )
            elif document in seen_docs:
                row_errors.append(
                    ("id_number", f"Documento {document} duplicado en el archivo (fila {seen_docs[document]}).")
                )
        if jersey is not None:
            if jersey in existing_jerseys:
                row_errors.append(
                    ("jersey_number", f"La camiseta #{jersey} ya está asignada en este equipo.")
                )
            elif jersey in seen_jerseys:
                row_errors.append(
                    ("jersey_number", f"Camiseta #{jersey} duplicada en el archivo (fila {seen_jerseys[jersey]}).")
                )

        if row_errors:
            for field, message in row_errors:
                errors.append({"row": r, "field": FIELD_LABELS[field], "message": message})
            continue

        seen_docs[document] = r
        seen_jerseys[jersey] = r
        to_create.append(
            Player(
                first_name=first_name[:100],
                last_name=last_name[:100],
                id_number=document[:50],
                jersey_number=jersey,
                position=position,
                birth_date=birth_date,
                phone=phone,
                team=team,
                tournament=team.tournament,
                posted_by=user,
            )
        )

    with transaction.atomic():
        if to_create:
            Player.objects.bulk_create(to_create, batch_size=DB_BATCH_SIZE)

    error_rows = len({e["row"] for e in errors})
    return {
        "total_rows": len(rows),
        "created": len(to_create),
        "error_count": error_rows,
        "errors": errors[:200],
        "message": _summary_message(len(to_create), error_rows),
    }
