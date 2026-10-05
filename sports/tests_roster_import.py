"""Carga masiva de la plantilla de jugadores desde Excel/CSV."""

import io
from datetime import timedelta

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from authentication.models import User
from organizations.models import Organization
from sports.models import Player, Team, Tournament
from sports.roster_import import ROSTER_HEADERS


def _xlsx(rows, headers=ROSTER_HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        "plantilla.xlsx",
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class RosterImportTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="Org Roster", slug="org-roster")
        self.owner = User.objects.create_user(
            email="owner-roster@test.com",
            username="owner_roster",
            password="SecurePass123!",
            organization=self.org,
            role="manager",
        )
        self.tournament = Tournament.objects.create(
            name="Copa Roster",
            slug="copa-roster",
            sport_type="football",
            organization=self.org,
            posted_by=self.owner,
            start_date="2026-09-01",
            end_date="2026-12-01",
        )
        self.team = Team.objects.create(
            name="Atlético Test",
            slug="atletico-test",
            abbreviation="ATL",
            tournament=self.tournament,
            organization=self.org,
            posted_by=self.owner,
            coach_email="coach@test.com",
        )
        Player.objects.create(
            first_name="Ya",
            last_name="Inscrito",
            id_number="555",
            jersey_number=7,
            position="forward",
            team=self.team,
            tournament=self.tournament,
            posted_by=self.owner,
        )
        self.level2 = User.objects.create_user(
            email="l2@test.com",
            username="l2_admin",
            password="SecurePass123!",
            organization=self.org,
            admin_level=2,
        )
        self.coach = User.objects.create_user(
            email="coach@test.com",
            username="coach_roster",
            password="SecurePass123!",
            organization=self.org,
            role="manager",
            sports_module_expires_at=timezone.now() + timedelta(days=30),
        )
        self.url = f"/api/v1/sports/teams/{self.team.id}/import-roster/"

    def _post(self, user, upload):
        client = APIClient()
        client.force_authenticate(user=user)
        return client.post(self.url, {"file": upload}, format="multipart")

    def test_level2_imports_and_gets_detailed_report(self):
        upload = _xlsx(
            [
                ["Juan", "Pérez", "1001", 10, "Delantero", "2001-05-14", "3001234567"],
                ["Luis", "Díaz", "1002", 1, "goalkeeper", "", ""],
                ["Mario", "Ruiz", "1003", 8, "Mediocampista", "", ""],
                ["Copia", "Doc", "1001", 11, "Defensa", "", ""],
                ["Copia", "Camiseta", "1004", 7, "Defensa", "", ""],
                ["Sin", "Posicion", "1005", 12, "Volador", "", ""],
            ]
        )
        res = self._post(self.level2, upload)
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["created"], 3)
        self.assertEqual(body["error_count"], 3)
        self.assertEqual(
            body["message"], "3 jugadores importados con éxito, 3 registros con errores"
        )
        self.assertEqual({e["row"] for e in body["errors"]}, {5, 6, 7})

        juan = Player.objects.get(team=self.team, id_number="1001")
        self.assertEqual(juan.position, "forward")
        self.assertEqual(juan.jersey_number, 10)
        self.assertEqual(str(juan.birth_date), "2001-05-14")
        self.assertEqual(juan.phone, "3001234567")
        self.assertEqual(juan.posted_by, self.level2)
        self.assertEqual(self.team.players.count(), 4)

    def test_superuser_can_import_by_slug(self):
        superuser = User.objects.create_superuser(
            email="root@test.com", username="root_roster", password="SecurePass123!"
        )
        client = APIClient()
        client.force_authenticate(user=superuser)
        res = client.post(
            f"/api/v1/sports/teams/{self.team.slug}/import-roster/?tournament={self.tournament.slug}",
            {"file": _xlsx([["Ana", "Mora", "2001", 3, "Defensa", "", ""]])},
            format="multipart",
        )
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()["created"], 1)

    def test_tournament_owner_can_import(self):
        self.owner.sports_module_expires_at = timezone.now() + timedelta(days=30)
        self.owner.save(update_fields=["sports_module_expires_at"])
        res = self._post(self.owner, _xlsx([["Ana", "Mora", "2001", 3, "Defensa", "", ""]]))
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()["created"], 1)

        client = APIClient()
        client.force_authenticate(user=self.owner)
        self.assertEqual(
            client.get("/api/v1/sports/teams/import-roster-template/").status_code, 200
        )

    def test_owner_of_another_tournament_is_forbidden(self):
        other = User.objects.create_user(
            email="other-owner@test.com",
            username="other_owner",
            password="SecurePass123!",
            organization=self.org,
            role="manager",
            sports_module_expires_at=timezone.now() + timedelta(days=30),
        )
        Tournament.objects.create(
            name="Otra Copa",
            slug="otra-copa",
            sport_type="football",
            organization=self.org,
            posted_by=other,
            start_date="2026-09-01",
            end_date="2026-12-01",
        )
        res = self._post(other, _xlsx([["Ana", "Mora", "2001", 3, "Defensa", "", ""]]))
        self.assertEqual(res.status_code, 403, res.content)
        self.assertEqual(self.team.players.count(), 1)

    def test_coach_without_super_admin_is_forbidden(self):
        res = self._post(self.coach, _xlsx([["Ana", "Mora", "2001", 3, "Defensa", "", ""]]))
        self.assertEqual(res.status_code, 403, res.content)
        self.assertEqual(self.team.players.count(), 1)

        client = APIClient()
        client.force_authenticate(user=self.coach)
        tpl = client.get("/api/v1/sports/teams/import-roster-template/")
        self.assertEqual(tpl.status_code, 403)

    def test_wrong_columns_are_rejected(self):
        upload = _xlsx([["Ana", "Mora"]], headers=["Nombre completo", "Equipo"])
        res = self._post(self.level2, upload)
        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn("Documento de Identidad", res.json()["missing_columns"])
        self.assertEqual(self.team.players.count(), 1)

    def test_template_download_has_expected_headers(self):
        client = APIClient()
        client.force_authenticate(user=self.level2)
        res = client.get("/api/v1/sports/teams/import-roster-template/")
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(io.BytesIO(res.content))
        headers = [c.value for c in wb.active[1]]
        self.assertEqual(headers, ROSTER_HEADERS)

    def test_csv_upload_is_supported(self):
        csv_content = (
            "Nombres,Apellidos,Documento de Identidad,Número de Camiseta,Posición\n"
            "Pedro,Lara,3001,22,Defensa\n"
        ).encode("utf-8-sig")
        upload = SimpleUploadedFile("plantilla.csv", csv_content, content_type="text/csv")
        res = self._post(self.level2, upload)
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()["created"], 1)
